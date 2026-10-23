# -*- coding: utf-8 -*-
"""Bundled script of the weekly-report-ppt skill: outline JSON → .pptx (python-pptx). 16:9, dark navy background.
make_skills.py copies this file into skills/weekly-report-ppt/scripts/."""
import json, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

INK, INDIGO, WHITE, SUB = RGBColor(0x11, 0x13, 0x18), RGBColor(0x63, 0x66, 0xF1), RGBColor(0xF5, 0xF6, 0xFA), RGBColor(0xA8, 0xAD, 0xBD)


def build(outline, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slides = outline.get("slides", [])
    for i, s in enumerate(slides):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = INK
        bar = sl.shapes.add_shape(1, Inches(0.6), Inches(0.55 if i else 3.0), Inches(0.12), Inches(0.9))
        bar.fill.solid(); bar.fill.fore_color.rgb = INDIGO; bar.line.fill.background()
        tb = sl.shapes.add_textbox(Inches(0.95), Inches(0.45 if i else 2.85), Inches(11.5), Inches(1.2))
        tb.name = "Title"
        p = tb.text_frame.paragraphs[0]; p.text = str(s.get("title", ""))
        p.font.size = Pt(40 if i == 0 else 32); p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = "Pretendard"
        body = sl.shapes.add_textbox(Inches(0.95), Inches(4.1 if i == 0 else 1.8), Inches(11.5), Inches(5))
        tf = body.text_frame; tf.word_wrap = True
        for j, b in enumerate(s.get("bullets", [])):
            q = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            q.text = ("" if i == 0 else "•  ") + str(b)
            q.font.size = Pt(22 if i == 0 else 24); q.font.color.rgb = SUB if i == 0 else WHITE; q.font.name = "Pretendard"
            q.space_after = Pt(14)
        if i:
            n = sl.shapes.add_textbox(Inches(12.2), Inches(6.85), Inches(0.8), Inches(0.4)).text_frame.paragraphs[0]
            n.text = str(i + 1); n.font.size = Pt(12); n.font.color.rgb = SUB
    prs.save(path)
    return len(slides)


if __name__ == "__main__":
    data = json.load(open(sys.argv[1], encoding="utf-8")) if len(sys.argv) > 1 else json.load(sys.stdin)
    out = sys.argv[2] if len(sys.argv) > 2 else data.get("file", "weekly.pptx")
    print(json.dumps({"ok": True, "path": out, "slides": build(data, out)}))
