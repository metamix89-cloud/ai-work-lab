# -*- coding: utf-8 -*-
"""EN ep03 lab — open skills only when needed (book ch.2 experiment 2-6 style), 9 Yummy Tummy Foods (fictional) work skills.
English re-run of the Korean ep038 lab. Real gpt-4.1-mini calls → results*.json. Graded automatically by rules written before the run.
  N  no skills      : the base assistant prompt only
  A  everything in  : all 9 SKILL.md files + reference files pasted into the system prompt
  B  on demand      : only name + description (when to use, when not to) → opens them with read_skill / read_skill_file
  D  vague catalog  : same as B, but every description blurred to 'helps with writing'
  B2 (T1 fix try)   : B + one line in the run_script description: 'read the SKILL.md first'
  B3 (T1 fix try)   : B + the code refuses run_script until that skill's SKILL.md has been read
Input tokens = usage.prompt_tokens summed over every round trip.
Usage:
  python make_skills.py                               (once)
  python lab.py N,A,B,D 5                             → results.json
  python lab.py B 10 results_t1.json --only=T1        → the PPT task alone
  python lab.py B2 5 results_t1fix.json --only=T1
  python lab.py B3 10 results_t1guard.json --only=T1"""
import json, os, re, sys, time
from concurrent.futures import ThreadPoolExecutor
from openai import OpenAI
from dotenv import load_dotenv
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
SK = os.path.join(HERE, "skills")
OUT = os.path.join(HERE, "out")
os.makedirs(OUT, exist_ok=True)
sys.path.insert(0, os.path.join(SK, "weekly-report-ppt", "scripts"))
import make_ppt

load_dotenv()   # OPENAI_API_KEY from a .env file in this folder or any parent folder
client = OpenAI()
MODEL = "gpt-4.1-mini"
TRIALS = 5

NAMES = sorted(os.listdir(SK))


def skill_md(n):
    with open(os.path.join(SK, n, "SKILL.md"), encoding="utf-8") as f:
        return f.read()


def desc(n):
    return re.search(r"^description: (.+)$", skill_md(n), re.M).group(1)


def extra_files(n):
    out = []
    for dp, _, fs in os.walk(os.path.join(SK, n)):
        for f in fs:
            rel = os.path.relpath(os.path.join(dp, f), os.path.join(SK, n)).replace("\\", "/")
            if rel != "SKILL.md" and rel.endswith(".md"):
                out.append(rel)
    return sorted(out)


VAGUE = {"weekly-report-ppt": "Helps write work documents.", "return-policy": "Helps with customer service.",
         "partner-email": "Helps with writing.", "meeting-minutes": "Helps organize documents.",
         "product-copy": "Helps with writing.", "inventory-alert": "Helps organize data.",
         "expense-report": "Helps with calculations.", "haccp-check": "Helps with work checks.",
         "onboarding-guide": "Helps with internal guides."}

BASE = ("You are a work assistant for employees of Yummy Tummy Foods (a fictional company). Answer in English. "
        "Use only the information you are given and never invent numbers.")


def system_prompt(cond):
    if cond == "N":
        return BASE
    if cond in ("B2", "B3"):
        cond = "B"
    if cond == "A":
        parts = [BASE, "\n# Company work manuals (if a manual matches the request, follow it exactly)\n"
                 "Run bundled scripts with run_script(skill, script, input).\n"]
        for n in NAMES:
            parts.append(f"\n===== skill: {n} =====\n{skill_md(n)}")
            for rel in extra_files(n):
                with open(os.path.join(SK, n, rel), encoding="utf-8") as f:
                    parts.append(f"\n----- {n}/{rel} -----\n" + f.read())
        return "".join(parts)
    d = desc if cond == "B" else (lambda n: VAGUE[n])
    cat = "\n".join(f"- {n}: {d(n)}" for n in NAMES)
    return (BASE + "\n\n# Installed skills (catalog only)\n"
            "If a skill matches the request, first read its manual with read_skill and follow it exactly. "
            "If the manual points to another file, read it with read_skill_file. If no skill matches, answer directly without reading any.\n" + cat)


def fn(name, d, props):
    return {"type": "function", "function": {"name": name, "description": d, "parameters": {
        "type": "object", "properties": {k: {"type": "string", "description": v} for k, v in props.items()},
        "required": list(props)}}}


T_READ = fn("read_skill", "Read a skill's full manual (SKILL.md).", {"name": "skill name"})
T_FILE = fn("read_skill_file", "Read a reference file inside a skill folder.", {"name": "skill name", "path": "path inside the skill folder, e.g. reference.md"})
T_RUN = fn("run_script", "Run a script bundled with a skill.",
           {"skill": "skill name", "script": "script path, e.g. scripts/make_ppt.py", "input": "JSON string passed to the script"})
T_RUN2 = fn("run_script", "Run a script bundled with a skill. Always read that skill's SKILL.md with read_skill first, and use the input format the manual specifies.",
            {"skill": "skill name", "script": "script path, e.g. scripts/make_ppt.py", "input": "JSON string passed to the script (format: see SKILL.md)"})


def tools(cond):
    return {"N": None, "A": [T_RUN], "B": [T_READ, T_FILE, T_RUN], "B2": [T_READ, T_FILE, T_RUN2],
            "B3": [T_READ, T_FILE, T_RUN], "D": [T_READ, T_FILE, T_RUN]}[cond]


def run_tool(name, a, log, tag):
    try:
        if name == "read_skill":
            log["read"].append(a["name"])
            return skill_md(a["name"])
        if name == "read_skill_file":
            p = a["path"].replace("\\", "/").lstrip("./")
            log["files"].append(f"{a['name']}/{p}")
            full = os.path.normpath(os.path.join(SK, a["name"], p))
            if not full.startswith(os.path.join(SK, a["name"])):
                return "error: outside the skill folder"
            with open(full, encoding="utf-8") as f:
                return f.read()
        if name == "run_script":
            if log.get("guard") and a.get("skill") not in log["read"]:      # B3: the order is enforced in code
                log["blocked"] = log.get("blocked", 0) + 1
                return f"refused: first read the manual with read_skill(name=\"{a.get('skill')}\"), then run again using the manual's input format."
            if a.get("skill") != "weekly-report-ppt" or "make_ppt" not in a.get("script", ""):
                return "error: no such script"
            data = a["input"] if isinstance(a["input"], dict) else json.loads(a["input"])
            log["script_input"] = str(a["input"])[:600]
            path = os.path.join(OUT, f"{tag}.pptx")
            n = make_ppt.build(data, path)
            log["pptx"] = path
            return json.dumps({"ok": True, "path": os.path.basename(path), "slides": n})
    except Exception as e:
        return f"error: {e}"
    return "error: unknown tool"


def agent(cond, msg, tag):
    messages = [{"role": "system", "content": system_prompt(cond)}, {"role": "user", "content": msg}]
    log = dict(read=[], files=[], pptx=None, prompt_tokens=0, completion_tokens=0, turns=0, first_prompt=0, guard=cond == "B3")
    t0 = time.time()
    for _ in range(8):
        kw = dict(model=MODEL, messages=messages)
        if tools(cond):
            kw["tools"] = tools(cond)
        r = client.chat.completions.create(**kw)
        log["turns"] += 1
        log["prompt_tokens"] += r.usage.prompt_tokens
        log["completion_tokens"] += r.usage.completion_tokens
        if not log["first_prompt"]:
            log["first_prompt"] = r.usage.prompt_tokens
        m = r.choices[0].message
        if not m.tool_calls:
            log["answer"] = m.content or ""
            break
        messages.append(m)
        for tc in m.tool_calls:
            try:
                a = json.loads(tc.function.arguments)
            except Exception:
                a = {}
            messages.append({"role": "tool", "tool_call_id": tc.id, "content": run_tool(tc.function.name, a, log, tag)})
    log.setdefault("answer", "")
    log["sec"] = round(time.time() - t0, 1)
    return log


# -- 10 requests (id, skill needed, request)
TASKS = [
    ("T1", "weekly-report-ppt", "Make this week's (10/13-10/17) sales report PPT.\n"
     "- Total sales $24,100 (target $23,000), last week $21,800\n- Pepperoni Pizza 2-pack $12,400 (target $11,000)\n"
     "- Mac & Cheese Bake $8,600 (target $9,000)\n- Cola 12-pack $3,100 (target $3,000)\n"
     "- Issue: the pizza sauce supplier announced a 3% price increase for November → asked another supplier for a quote\n"
     "- Next week: holiday promo talks with a big-box retailer (Alex), mac & cheese stock check (Morgan)"),
    ("T2", "return-policy", "Decide this customer request: 'I got the Yummy Tummy Handmade Burger Patties 5 days ago and haven't opened them. "
     "I changed my mind and want to return them. I paid $23.90.'"),
    ("T3", "partner-email", "Our supplier Sunrise Supply (contact: Jenna Lee) sent an email titled 'November price adjustment notice'. "
     "They're raising the pizza sauce price by 3% from November 1. Write a reply saying we'll review it internally and answer by Tuesday, October 28."),
    ("T4", "meeting-minutes", "Turn these notes into meeting minutes.\n10/20 new product meeting. Attendees: Alex (lead), Taylor, Morgan. "
     "Sesame Burger Buns launch confirmed for the 2nd week of November. Price set at $6.90. Taylor does the online store page by 10/31. "
     "Morgan does the in-store sign draft, no deadline set. Packaging design option B came up but no conclusion."),
    ("T5", "product-copy", "Write copy for the new product 'Yummy Tummy Sesame Burger Buns 8-pack'. Ingredients: wheat flour, water, sesame seeds, sugar, "
     "soybean flour, yeast, salt. Soft buns for burgers and sliders."),
    ("T6", "inventory-alert", "Here's the warehouse stock sheet. Write the alert.\nPepperoni Pizza 2-pack: stock 300, weekly average sales 260\n"
     "Mac & Cheese Bake: stock 500, weekly average sales 200\nCola 12-pack: stock 140, weekly average sales 100"),
    ("T7", "expense-report", "Calculate my Chicago trip reimbursement. Tue 10/14 - Thu 10/16, 2 nights / 3 days.\n- Round-trip train $119.60 (receipt)\n"
     "- Hotel 2 nights, $85.00 per night (receipt)\n- 6 meals, $12.00 each (receipts)\n- Taxi $15.00 (no receipt)"),
    ("T8", None, "Explain what Excel's VLOOKUP does in two lines."),
    ("T9", None, "Write a message for the team group chat announcing Friday's team dinner: 7 pm, the BBQ place in front of the office."),
    ("T10", "return-policy", "Decide this customer request: 'I opened the BBQ Sauce 18 oz and tried it, but it smells sour. I think it's gone bad. I paid $8.50.'"),
]


def first_line(t):
    return next((l.strip() for l in t.splitlines() if l.strip()), "")


def line_with(t, key):
    key = key.lower()
    return next((l for l in t.splitlines() if key in l.lower()), "")


def grade_text(tid, t, read):
    """Grades every task except T1 (the PPT file). Returns (ok, note)."""
    if tid in ("T2", "T10"):
        want = "[Decision] Not returnable" if tid == "T2" else "[Decision] Returnable"
        right = want in t and (tid == "T2" or "8.50" in t)
        fmt = first_line(t).strip("*` ").startswith(want)
        return right and fmt, ("right + format OK" if right and fmt else "right, first-line format broken" if right else "wrong decision")
    if tid == "T3":
        ok = ("Yummy Tummy Foods Sales Support Team" in t and "[Yummy Tummy Foods] Re:" in t
              and not re.search(r"\bwe (accept|agree)\b", t, re.I))
        return ok, "signature/subject " + ("present" if ok else "missing")
    if tid == "T4":
        ok = all(h in t for h in ["■ Decisions", "■ Action Items", "■ Next Meeting"]) and "due tbd" in t.lower()
        return ok, "standard headers " + ("present" if ok else "missing")
    if tid == "T5":
        l = line_with(t, "[Allergens]").lower()
        ok = bool(l) and "soy" in l and "wheat" in l and "sesame" in l
        return ok, l.strip()[:50] or "no allergen line"
    if tid == "T6":
        # Grade the alert line (the one with a 🔴/🟢/⚪ marker), not the first line that names the item:
        # models often write their calculation first. (Bug in the first grading — see NOTES.md.)
        def alert(k):
            return next((l for l in t.splitlines() if k.lower() in l.lower() and any(e in l for e in "🔴🟢⚪")), "")
        ok = "🔴" in alert("Pizza") and "🟢" in alert("Mac & Cheese") and "🔴" in alert("Cola")
        em = lambda l: "🔴" if "🔴" in l else "🟢" if "🟢" in l else "-"
        return ok, " / ".join(em(alert(k)) for k in ["Pizza", "Mac & Cheese", "Cola"])
    if tid == "T7":
        l = line_with(t, "Total reimbursement:")
        return "379.60" in l, l.strip()[:40] or "no total line"
    if tid == "T8":
        return not read, "skills read: " + (",".join(read) or "none")
    if tid == "T9":
        ok = "partner-email" not in read and "Sales Support Team" not in t and "[Yummy Tummy Foods] Re" not in t
        return ok, "skills read: " + (",".join(read) or "none")


def grade_pptx(path):
    if not path:
        return False, "no PPT file"
    prs = Presentation(path)
    titles = [next((s.text_frame.text for s in sl.shapes if s.name == "Title"), "") for sl in prs.slides]
    alltext = " ".join(s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
    if not titles:
        return False, "empty PPT (0 slides)"
    ok = (len(titles) == 5 and titles[0].startswith("Weekly Sales Report |") and titles[-1].strip() == "Next Week To-Do"
          and bool(re.search(r"\$\d+(\.\d)?K", alltext)))
    return ok, f"{len(titles)} slides · {titles[0][:24]} … {titles[-1][:16]}"


def grade(tid, log):
    if tid == "T1":
        return grade_pptx(log["pptx"])
    return grade_text(tid, log["answer"], log["read"])


def one(job):
    cond, (tid, need, msg), seed = job
    tag = f"{cond}_{tid}_{seed}"
    for k in range(3):
        try:
            log = agent(cond, msg, tag)
            break
        except Exception as e:
            if k == 2:
                return dict(cond=cond, task=tid, seed=seed, error=str(e))
            time.sleep(5)
    ok, note = grade(tid, log)
    return dict(cond=cond, task=tid, seed=seed, need=need, ok=ok, note=note, read=log["read"], files=log["files"],
                prompt_tokens=log["prompt_tokens"], first_prompt=log["first_prompt"], completion_tokens=log["completion_tokens"],
                turns=log["turns"], sec=log["sec"], script_input=log.get("script_input"), blocked=log.get("blocked", 0), answer=log["answer"][:1500])


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--only=")]
    only = next((a.split("=", 1)[1].split(",") for a in sys.argv[1:] if a.startswith("--only=")), None)
    conds = args[0].split(",") if args else ["N", "A", "B", "D"]
    trials = int(args[1]) if len(args) > 1 else TRIALS
    jobs = [(c, t, s) for c in conds for t in TASKS if not only or t[0] in only for s in range(trials)]
    with ThreadPoolExecutor(8) as ex:
        res = list(ex.map(one, jobs))
    fn_out = os.path.join(HERE, args[2] if len(args) > 2 else "results.json")
    with open(fn_out, "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=1)
    for c in conds:
        rs = [r for r in res if r["cond"] == c and "error" not in r]
        print(c, "correct", sum(r["ok"] for r in rs), "/", len(rs), "· avg input tokens", round(sum(r["prompt_tokens"] for r in rs) / max(1, len(rs))),
              "· first call", round(sum(r["first_prompt"] for r in rs) / max(1, len(rs))))
