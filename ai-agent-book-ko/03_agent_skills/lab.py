# -*- coding: utf-8 -*-
"""ep038 실습 — 스킬을 필요할 때만 펼치기 (원서 2장 실험 2-6 방식, 가나다식품 업무 스킬 9개로 각색).
조건 4개 × 요청 10건 × 3회 실제 실행 → results.json. 채점은 사람이 미리 정한 정답 규칙으로 자동.
  N 스킬 없음   : 업무 비서 기본 프롬프트만
  A 전부 넣기   : 스킬 9개 설명서 + 참고 파일 전체를 시스템 프롬프트에
  B 필요할 때만 : 이름 + 설명(언제 쓰고 언제 안 쓰는지)만 → read_skill · read_skill_file 로 펼침
  D 설명 대충   : B 와 같은 구조, 설명만 '글쓰기를 돕는다' 식으로 뭉뚱그림
입력 토큰은 API 가 돌려준 usage.prompt_tokens 를 모든 왕복에 걸쳐 더한 값."""
import json, os, re, sys, time
from concurrent.futures import ThreadPoolExecutor
from openai import OpenAI
from dotenv import load_dotenv
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SK = os.path.join(HERE, "skills")
OUT = os.path.join(HERE, "out"); os.makedirs(OUT, exist_ok=True)
sys.path.insert(0, os.path.join(SK, "weekly-report-ppt", "scripts"))
import make_ppt

load_dotenv()
client = OpenAI()
MODEL = "gpt-4.1-mini"
TRIALS = 3

NAMES = sorted(os.listdir(SK))


def skill_md(n):
    return open(os.path.join(SK, n, "SKILL.md"), encoding="utf-8").read()


def desc(n):
    return re.search(r"^description: (.+)$", skill_md(n), re.M).group(1)


def extra_files(n):
    out = []
    for dp, _, fs in os.walk(os.path.join(SK, n)):
        for f in fs:
            rel = os.path.relpath(os.path.join(dp, f), os.path.join(SK, n)).replace("\\", "/")
            if rel != "SKILL.md" and rel.endswith(".md"):
                out.append(rel)
    return sorted(out)


VAGUE = {"weekly-report-ppt": "업무 문서 작성을 돕는다.", "return-policy": "고객 응대를 돕는다.",
         "partner-email": "글쓰기를 돕는다.", "meeting-minutes": "문서 정리를 돕는다.",
         "product-copy": "글쓰기를 돕는다.", "inventory-alert": "데이터 정리를 돕는다.",
         "expense-report": "숫자 계산을 돕는다.", "haccp-check": "업무 점검을 돕는다.",
         "onboarding-guide": "사내 안내를 돕는다."}

BASE = ("너는 가나다식품(가상 회사) 직원들을 돕는 업무 비서다. 한국어로 답한다. "
        "받은 정보만 쓰고 모르는 숫자는 지어내지 않는다.")


def system_prompt(cond):
    if cond == "N":
        return BASE
    if cond in ("B2", "B3"):   # B3: B 와 같고, 설명서를 안 읽고 실행하면 코드가 돌려보냄
        cond = "B"
    if cond == "A":
        parts = [BASE, "\n# 회사 업무 설명서 (요청에 해당하는 설명서가 있으면 그대로 따른다)\n"
                 "번들 스크립트는 run_script(skill, script, input) 로 실행한다.\n"]
        for n in NAMES:
            parts.append(f"\n===== 스킬: {n} =====\n{skill_md(n)}")
            for rel in extra_files(n):
                parts.append(f"\n----- {n}/{rel} -----\n" + open(os.path.join(SK, n, rel), encoding="utf-8").read())
        return "".join(parts)
    d = desc if cond == "B" else (lambda n: VAGUE[n])
    cat = "\n".join(f"- {n}: {d(n)}" for n in NAMES)
    return (BASE + "\n\n# 설치된 스킬 (목록만 있음)\n"
            "요청에 맞는 스킬이 있으면 먼저 read_skill 로 설명서를 읽고 그대로 따른다. "
            "설명서가 다른 파일을 가리키면 read_skill_file 로 읽는다. 맞는 스킬이 없으면 읽지 않고 바로 답한다.\n" + cat)


def fn(name, d, props):
    return {"type": "function", "function": {"name": name, "description": d, "parameters": {
        "type": "object", "properties": {k: {"type": "string", "description": v} for k, v in props.items()},
        "required": list(props)}}}


T_READ = fn("read_skill", "스킬 설명서(SKILL.md) 전체를 읽는다.", {"name": "스킬 이름"})
T_FILE = fn("read_skill_file", "스킬 폴더 안의 참고 파일을 읽는다.", {"name": "스킬 이름", "path": "스킬 폴더 기준 경로. 예: reference.md"})
T_RUN = fn("run_script", "스킬에 들어 있는 스크립트를 실행한다.",
           {"skill": "스킬 이름", "script": "스크립트 경로. 예: scripts/make_ppt.py", "input": "스크립트에 넘길 JSON 문자열"})


T_RUN2 = fn("run_script", "스킬에 들어 있는 스크립트를 실행한다. 반드시 그 스킬의 SKILL.md 를 read_skill 로 읽고, 설명서가 정한 입력 형식대로 실행한다.",
            {"skill": "스킬 이름", "script": "스크립트 경로. 예: scripts/make_ppt.py", "input": "스크립트에 넘길 JSON 문자열 (형식은 SKILL.md 참고)"})


def tools(cond):
    # B2: B 와 같고 실행 도구 설명에 '설명서 먼저 읽기' 한 줄만 추가
    return {"N": None, "A": [T_RUN], "B": [T_READ, T_FILE, T_RUN], "B2": [T_READ, T_FILE, T_RUN2], "B3": [T_READ, T_FILE, T_RUN], "D": [T_READ, T_FILE, T_RUN]}[cond]


def run_tool(name, a, log, tag):
    try:
        if name == "read_skill":
            log["read"].append(a["name"]); return skill_md(a["name"])
        if name == "read_skill_file":
            p = a["path"].replace("\\", "/").lstrip("./")
            log["files"].append(f"{a['name']}/{p}")
            full = os.path.normpath(os.path.join(SK, a["name"], p))
            if not full.startswith(os.path.join(SK, a["name"])):
                return "오류: 스킬 폴더 밖"
            return open(full, encoding="utf-8").read()
        if name == "run_script":
            if log.get("guard") and a.get("skill") not in log["read"]:
                log["blocked"] = log.get("blocked", 0) + 1
                return f"실행 거부: 먼저 read_skill(name=\"{a.get('skill')}\") 로 설명서를 읽고, 설명서의 입력 형식대로 다시 실행하세요."
            if a.get("skill") != "weekly-report-ppt" or "make_ppt" not in a.get("script", ""):
                return "오류: 그런 스크립트 없음"
            data = a["input"] if isinstance(a["input"], dict) else json.loads(a["input"])
            log["script_input"] = str(a["input"])[:600]
            path = os.path.join(OUT, f"{tag}.pptx")
            n = make_ppt.build(data, path)
            log["pptx"] = path
            return json.dumps({"ok": True, "path": os.path.basename(path), "slides": n}, ensure_ascii=False)
    except Exception as e:
        return f"오류: {e}"
    return "오류: 알 수 없는 도구"


def agent(cond, msg, tag):
    messages = [{"role": "system", "content": system_prompt(cond)}, {"role": "user", "content": msg}]
    log = dict(read=[], files=[], pptx=None, prompt_tokens=0, completion_tokens=0, turns=0, first_prompt=0, guard=cond == "B3")
    t0 = time.time()
    for _ in range(8):
        kw = dict(model=MODEL, messages=messages)
        if tools(cond):
            kw["tools"] = tools(cond)
        r = client.chat.completions.create(**kw)
        log["turns"] += 1
        log["prompt_tokens"] += r.usage.prompt_tokens; log["completion_tokens"] += r.usage.completion_tokens
        if not log["first_prompt"]:
            log["first_prompt"] = r.usage.prompt_tokens
        m = r.choices[0].message
        if not m.tool_calls:
            log["answer"] = m.content or ""; break
        messages.append(m)
        for tc in m.tool_calls:
            try:
                a = json.loads(tc.function.arguments)
            except Exception:
                a = {}
            messages.append({"role": "tool", "tool_call_id": tc.id, "content": run_tool(tc.function.name, a, log, tag)})
    log.setdefault("answer", "")
    log["sec"] = round(time.time() - t0, 1)
    return log


# ── 요청 10건 (id, 필요한 스킬, 요청 문장)
TASKS = [
    ("T1", "weekly-report-ppt", "이번 주(10/13~10/17) 영업 보고 PPT 만들어 줘.\n"
     "- 총매출 2,410만원 (목표 2,300만원), 지난주 2,180만원\n- 냉동만두 1,240만원 (목표 1,100만원)\n"
     "- 떡볶이 밀키트 860만원 (목표 900만원)\n- 참기름 310만원 (목표 300만원)\n"
     "- 이슈: 떡볶이 소스 공급사가 11월 단가 3% 인상 통보 → 대체 공급사 견적 요청\n"
     "- 다음 주: 대형마트 연말 행사 협의(김 대리), 밀키트 재고 점검(박 주임)"),
    ("T2", "return-policy", "고객 문의 판정해 줘: '가나다 수제 떡갈비 받은 지 5일 됐는데 안 뜯었어요. 마음이 바뀌어서 반품하고 싶어요. 결제 23,900원이에요.'"),
    ("T3", "partner-email", "거래처 가나공급 이수진 과장이 '11월 납품 단가 조정 안내' 메일을 보냈어. 떡볶이 소스 단가를 11월 1일부터 3% 올린다는 내용이야. "
     "내부 검토 후 10월 28일(화)까지 답을 주겠다고 답장 써 줘."),
    ("T4", "meeting-minutes", "아래 메모를 회의록으로 정리해 줘.\n10/20 신제품 회의. 참석 김팀장 이대리 박주임. 들기름 김자반 11월 둘째 주 출시 확정. "
     "가격 6,900원으로 결정. 이대리가 온라인몰 상세페이지 10/31까지. 박주임 매장 POP 시안, 기한은 안 정함. 포장 디자인 B안 얘기 나왔는데 결론 못 냄."),
    ("T5", "product-copy", "신제품 '가나다 들기름 김자반' 소개 문구 써 줘. 원재료: 김(국산), 들기름, 참깨, 간장, 설탕. 밥반찬, 주먹밥용."),
    ("T6", "inventory-alert", "물류센터 재고표야. 알림 써 줘.\n냉동만두 1kg: 재고 300, 주평균 판매 260\n떡볶이 밀키트: 재고 500, 주평균 판매 200\n참기름 350ml: 재고 140, 주평균 판매 100"),
    ("T7", "expense-report", "부산 출장 정산해 줘. 10/14(화)~10/16(목) 2박 3일.\n- KTX 왕복 119,600원 (영수증 있음)\n"
     "- 숙박 2박, 1박 85,000원 (영수증 있음)\n- 식사 6끼, 끼니마다 12,000원 (영수증 있음)\n- 택시 15,000원 (영수증 없음)"),
    ("T8", None, "엑셀 VLOOKUP 함수가 뭔지 두 줄로 알려 줘."),
    ("T9", None, "팀 단톡방에 올릴 금요일 회식 공지 문자 써 줘. 저녁 7시, 회사 앞 고깃집."),
    ("T10", "return-policy", "고객 문의 판정해 줘: '고추장 500g 뜯어서 먹어 봤는데 시큼한 냄새가 나요. 변질된 것 같아요. 결제 8,500원.'"),
]


def first_line(t):
    return next((l.strip() for l in t.splitlines() if l.strip()), "")


def line_with(t, key):
    return next((l for l in t.splitlines() if key in l), "")


def grade(tid, log):
    t = log["answer"]
    if tid == "T1":
        if not log["pptx"]:
            return False, "PPT 파일 없음"
        prs = Presentation(log["pptx"])
        titles = [next((s.text_frame.text for s in sl.shapes if s.name == "Title"), "") for sl in prs.slides]
        alltext = " ".join(s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame)
        if not titles:
            return False, "빈 PPT (0장)"
        ok = len(titles) == 5 and titles[0].startswith("주간 영업 보고 |") and titles[-1].strip() == "다음 주 할 일" and "만원" in alltext
        return ok, f"{len(titles)}장 · {titles[0][:20]} … {titles[-1][:12] if titles else ''}"
    if tid in ("T2", "T10"):
        want = "[판정] 반품 불가" if tid == "T2" else "[판정] 반품 가능"
        right = want in t and (tid == "T2" or "8,500" in t)
        fmt = want in first_line(t)
        return right and fmt, ("정답·형식 OK" if right and fmt else "판정은 맞음, 첫 줄 형식 어김" if right else "판정 틀림")
    if tid == "T3":
        ok = "가나다식품 영업지원팀 드림" in t and "[가나다식품] Re:" in t and not re.search(r"수락합니다|동의합니다", t)
        return ok, "서명·제목 " + ("있음" if ok else "없음")
    if tid == "T4":
        ok = all(h in t for h in ["■ 결정 사항", "■ 할 일", "■ 다음 회의"]) and "기한 미정" in t
        return ok, "표준 머리표 " + ("있음" if ok else "없음")
    if tid == "T5":
        l = line_with(t, "[알레르기]")
        return bool(l) and "대두" in l and "밀" in l, l.strip()[:40] or "알레르기 줄 없음"
    if tid == "T6":
        ok = "🔴" in line_with(t, "냉동만두") and "🟢" in line_with(t, "떡볶이") and "🔴" in line_with(t, "참기름")
        em = lambda l: "🔴" if "🔴" in l else "🟢" if "🟢" in l else "-"
        return ok, " / ".join(em(line_with(t, k)) for k in ["냉동만두", "떡볶이", "참기름"])
    if tid == "T7":
        l = line_with(t, "정산액:")
        return "379,600" in l, l.strip()[:30] or "정산액 줄 없음"
    if tid == "T8":
        return not log["read"], "읽은 스킬: " + (",".join(log["read"]) or "없음")
    if tid == "T9":
        ok = "partner-email" not in log["read"] and "영업지원팀 드림" not in t and "[가나다식품] Re" not in t
        return ok, "읽은 스킬: " + (",".join(log["read"]) or "없음")


def one(job):
    cond, (tid, need, msg), seed = job
    tag = f"{cond}_{tid}_{seed}"
    for k in range(3):
        try:
            log = agent(cond, msg, tag); break
        except Exception as e:
            if k == 2:
                return dict(cond=cond, task=tid, seed=seed, error=str(e))
            time.sleep(5)
    ok, note = grade(tid, log)
    return dict(cond=cond, task=tid, seed=seed, need=need, ok=ok, note=note, read=log["read"], files=log["files"],
                prompt_tokens=log["prompt_tokens"], first_prompt=log["first_prompt"], completion_tokens=log["completion_tokens"],
                turns=log["turns"], sec=log["sec"], script_input=log.get("script_input"), blocked=log.get("blocked", 0), answer=log["answer"][:1500])


if __name__ == "__main__":
    conds = sys.argv[1].split(",") if len(sys.argv) > 1 else ["N", "A", "B", "D"]
    trials = int(sys.argv[2]) if len(sys.argv) > 2 else TRIALS
    only = os.environ.get("ONLY")
    jobs = [(c, t, s) for c in conds for t in TASKS if not only or t[0] in only.split(",") for s in range(trials)]
    with ThreadPoolExecutor(8) as ex:
        res = list(ex.map(one, jobs))
    fn_out = os.path.join(HERE, sys.argv[3] if len(sys.argv) > 3 else "results.json")
    json.dump(res, open(fn_out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for c in conds:
        rs = [r for r in res if r["cond"] == c and "error" not in r]
        print(c, "정답", sum(r["ok"] for r in rs), "/", len(rs), "· 평균 입력 토큰", round(sum(r["prompt_tokens"] for r in rs) / max(1, len(rs))),
              "· 첫 요청", round(sum(r["first_prompt"] for r in rs) / max(1, len(rs))))
